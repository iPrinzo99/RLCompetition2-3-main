from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy
D=Document(); sec=D.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.68);sec.left_margin=sec.right_margin=Inches(.75)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=D.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7)
D.styles['Normal'].font.size=Pt(10.5);D.styles['Normal'].paragraph_format.line_spacing=1.08
D.styles['Title'].font.size=Pt(25);D.styles['Heading 1'].font.size=Pt(16);D.styles['Heading 2'].font.size=Pt(12)
D.styles['Heading 1'].paragraph_format.space_before=Pt(7)
D.styles['Heading 2'].paragraph_format.space_before=Pt(8)
sec.header.paragraphs[0].text='CIS 5930    |    Competitions 2 and 3'
sec.header.paragraphs[0].runs[0].font.size=Pt(9)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Icy Grid World    •    ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
D.core_properties.title='Comparing Reinforcement Learning Policies in Icy Grid World'
D.core_properties.subject='Competitions 2 and 3 using the exact Competition 1 MDP'
D.core_properties.author=''
def p(s): return D.add_paragraph(s)
def h(s):D.add_heading(s,2)
def page(s):
 pp=D.add_heading(s,1);pp.paragraph_format.page_break_before=True
def mr(s):
 r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);return r
def sub(a,b):
 e=OxmlElement('m:sSub');ba=OxmlElement('m:e');ba.append(mr(a) if isinstance(a,str) else a);su=OxmlElement('m:sub');su.append(mr(b));e.extend([ba,su]);return e
def sup(a,b):
 e=OxmlElement('m:sSup');ba=OxmlElement('m:e');ba.append(mr(a) if isinstance(a,str) else a);su=OxmlElement('m:sup');su.append(mr(b));e.extend([ba,su]);return e
def summation(lo,hi):
 e=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
 ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch)
 loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'subSup');pr.append(loc);e.append(pr)
 for tag,txt in [('sub',lo),('sup',hi)]:
  z=OxmlElement('m:'+tag);z.append(mr(txt));e.append(z)
 e.append(OxmlElement('m:e'));return e
def mathparts(x):
 if not isinstance(x,str):return [x]
 for token,key in [('maxᵦ','b'),('maxₐ','a')]:
  if token in x:
   a,b=x.split(token,1);return mathparts(a)+[sub('max',key)]+mathparts(b)
 return [mr(x)]
def eq(*parts):
 pp=D.add_paragraph();pp.paragraph_format.space_after=Pt(8)
 om=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath')
 nodes=[]
 for x in parts:nodes.extend(mathparts(x))
 i=0
 while i<len(nodes):
  node=nodes[i]
  if node.tag==qn('m:nary') and i+1<len(nodes):
   node.find(qn('m:e')).append(nodes[i+1]);i+=1
  math.append(node);i+=1
 om.append(math);pp._p.append(om)
def table(headers,rows,widths):
 t=D.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cs=t.add_row().cells
  for i,x in enumerate(row):cs[i].text=x
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trPr.append(no)
  if ri==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F4F6F8' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.line_spacing=1.03
    for r in pp.runs:r.font.size=Pt(9.5);r.bold=ri==0
 D.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
from pathlib import Path
import json,numpy as np
ROOT=Path(__file__).parent
r=json.loads((ROOT/'results/results.json').read_text());diag=json.loads((ROOT/'results/diagnostics.json').read_text());seeds=json.loads((ROOT/'figures/representative_seeds.json').read_text())
summary={(x['algorithm'],x['exploration']):x for x in r['summary']}
names={'sarsa':'SARSA','q_learning':'Q-learning'}
def f(x,n=3):return f'{x:.{n}f}'
def mean_sd(x,key,n=3):return f(x[key+'_mean'],n)+' ± '+f(x[key+'_sd'],n)
D.core_properties.title='Optimal Policy Investigation for Competitions 2 and 3'
D.add_paragraph('Optimal Policy Investigation in Icy Grid World','Title')
D.add_paragraph('Competitions 2 and 3','Subtitle')
p('Policy iteration and value iteration produce the same optimal policy on the exact Competition 1 MDP. Their validated optimal start value is 26.282822174. The learned SARSA and Q-learning policies reach the goal reliably in held-out rollouts, but retain measurable discounted value loss after one million training transitions per run. They should be described as learned policies, rather than experimentally proven optimal policies.')
p('The original Competition 1 environment is preserved. The original folder and Red Team Solution archive contain identical environment source files; the four control algorithms and the experiments in this report are new additions. The conclusions below distinguish exact model calculations, sampled evaluation results, and convergence theory.')
D.add_heading('1 Verified comparison',1)
rows=[('Policy iteration','None','26.283','100.0%'),('Value iteration','None','26.283','100.0%')]
for mode in ['fixed','decay']:
 for alg in ['sarsa','q_learning']:
  x=summary[alg,mode];rows.append((names[alg],'0.10 fixed' if mode=='fixed' else '1.00 to 0.02',mean_sd(x,'value_start'),f(x['optimal_actions_fraction_mean']*100,1)+'%'))
table(['Algorithm','Training ε','Exact greedy start value','Optimal action fraction'],rows,[1.25,1.3,2.55,1.9])
p('Learning rows show mean ± sample standard deviation across ten independent training seeds. Values are obtained by exact evaluation of each frozen greedy policy on the underlying MDP. The optimal action fraction uses 351 nonterminal states and accepts actions within 10⁻⁶ reward units of the reference optimum. It is an all-state measure, not a success rate.')
h('What the measured differences mean')
p('Decaying-exploration SARSA has the highest mean learned-policy value in these settings. This is a descriptive result: approximate paired 95% intervals for SARSA minus Q-learning include zero under both fixed and decaying exploration. The experiment does not establish a universal winner. Planning has access to the full model, whereas learning uses sampled transitions.')
h('The required environment')
p('The fixed map is 20 × 20 with 48 impassable obstacles and 352 traversable states, including the goal. Start is (19, 0), goal is (0, 19), γ=0.99, and there are no hazards. Requested actions execute as intended with probability 0.80 and slip in each perpendicular direction with probability 0.10. Goal entry receives +100; every other nonterminal transition receives −1. The 500-step TimeLimit is an external cutoff [1].')

page('2 Exact MDP and evaluation objective')
p('Rows increase downward and columns increase rightward. The obstacle set O is the union of the following five sets; S is the 20 × 20 coordinate grid with O removed [1].')
for idx,body in [('1','{(r, 5) : 0 ≤ r ≤ 15, r ∉ {4, 12}}'),('2','{(r, 10) : 4 ≤ r ≤ 19, r ∉ {8, 16}}'),('3','{(r, 15) : 0 ≤ r ≤ 15, r ∉ {6, 13}}'),('4','{(9, c) : 6 ≤ c ≤ 9, c ≠ 8}'),('5','{(14, c) : 11 ≤ c ≤ 14, c ≠ 13}')]:eq(sub('O',idx),' = '+body)
p('Actions are numbered 0 up, 1 right, 2 down, and 3 left. Let m(s,b) be the cell reached in executed direction b, or s when a boundary or rock blocks the move. Collisions consume a step and receive the ordinary −1 reward. For s≠g:')
eq('P(s′ | s,a) = 0.80 𝟙{s′ = m(s,a)}')
eq('                + 0.10 𝟙{s′ = m(s,(a−1) mod 4)}')
eq('                + 0.10 𝟙{s′ = m(s,(a+1) mod 4)}')
p('Merge probabilities for duplicate successors. The goal is the only terminal state and is absorbing in the mathematical model, with zero reward thereafter. The +100 goal reward replaces the −1 step reward; it is not +99.')
eq('R(s,a) = 100 P(g | s,a) − [1 − P(g | s,a)]     for s ≠ g')
h('Discounted optimality and capped performance')
eq(sup('V','π'),'(s) = ',sup('E','π'),'[',summation('t=0','∞'),sup('γ','t'),sub('R','t+1'),' | ',sub('S','0'),' = s],     γ = 0.99')
p('For first goal arrival on transition T, the discounted trajectory return is:')
eq('G(T) = −100 + 200 · ',sup('0.99','T−1'))
p('Earlier arrival improves the return of that trajectory, but expected discounted return depends on the full arrival-time distribution, not only expected path length.')
p('The 500-step success rate and capped return are separate rollout metrics. No extra time coordinate is added to this stationary MDP. Use observation 20r+c: start ID 380 and goal ID 19. The declared Discrete(400) observation space includes 48 unused obstacle IDs, which must be excluded from planning.')

page('3 Algorithms and correct learning targets')
h('Policy iteration and value iteration')
p('Policy iteration evaluates a deterministic policy by solving the Bellman equations, then improves the action at each state. The implementation starts with action 0 and retains the current action on a tie [3].')
eq(sup('V','π'),'(s) = R(s,π(s)) + γ ',sub('∑','s′∈S'),'P(s′ | s,π(s)) ',sup('V','π'),'(s′)')
eq(sub('π','new'),'(s) ∈ arg maxₐ [R(s,a) + γ ',sub('∑','s′∈S'),'P(s′ | s,a) ',sup('V','π'),'(s′)]')
p('Value iteration starts at zero and applies Bellman optimality backups until the maximum change is below 10⁻¹⁰ [4]. Final extraction uses the smallest action ID among actions tied within 10⁻¹⁰.')
eq(sub('V','k+1'),'(s) = maxₐ [R(s,a) + γ ',sub('∑','s′∈S'),'P(s′ | s,a) ',sub('V','k'),'(s′)]')
p('Both methods keep V(g)=0. In the recorded run, policy iteration used 10 evaluation and improvement rounds; value iteration used 108 sweeps. Their maximum value difference was about 7.1 × 10⁻¹¹, and their selected actions agreed at all 351 nonterminal states. The optimality residuals were 2.84 × 10⁻¹⁴ and 4.50 × 10⁻¹¹, respectively.')
h('SARSA and Q learning')
p('Both tabular learners initialize Q to zero and use constant α=0.10. Let d=1 only for true goal termination. SARSA uses the next action sampled from its behavior policy; Q-learning uses the largest next-state Q estimate [5].')
eq('Q(s,a) ← Q(s,a) + α [y − Q(s,a)]')
eq(sub('y','SARSA'),' = r + γ(1−d) Q(s′,a′)')
eq(sub('y','Q'),' = r + γ(1−d) maxᵦ Q(s′,b)')
p('On a nonterminal truncation, keep the bootstrap term. SARSA samples a′ at the final pre-reset observation. Both flags end the episode and trigger reset, but only terminated=True suppresses continuation. Reaching the goal on step 500 can set both flags [1, 6].')
h('Theoretical expectation')
p('The finite discounted MDP has a unique optimal value function, although multiple optimal actions can exist. Standard tabular guarantees require adequate state–action coverage and suitable diminishing step sizes; SARSA also requires greedy-in-the-limit behavior with infinite exploration for convergence to Q*. Q-learning can target Q* under continued exploratory behavior. Constant α and the ε floor used here do not establish those asymptotic conditions [7].')
p('Q-learning remains sensitive to environmental slipping: its maximum chooses a requested action, not an executed direction. Neither learner can remove the 0.10/0.10 slips by choosing greedily.')

page('4 Experimental methodology and measured performance')
p('Each learner was trained with environment seeds 0–9 under two exploration settings, using 1,000,000 actual Gymnasium transitions per seed. The action RNG used seed+100000. Fixed exploration uses ε=0.10. The decay setting decreases ε linearly from 1.00 to 0.02 during the first 800,000 transitions, then holds 0.02. Random exploration selects uniformly among all four requested actions. Total learning interactions were 40 million.')
p('For each frozen greedy policy, 1,000 fresh episodes used evaluation seeds 200000–200999, always from the fixed start, with exploration disabled and ice retained. The two planning policies were also evaluated. This yields 42,000 policy rollouts; shared seeds make cross-policy comparisons paired rather than independent.')
rows=[]
for mode in ['fixed','decay']:
 for alg in ['sarsa','q_learning']:
  x=summary[alg,mode];rows.append((names[alg]+' '+mode,mean_sd(x,'start_gap'),mean_sd(x,'rollout_mean_steps'),mean_sd(x,'rollout_mean_capped_return')))
table(['Learning setting','Exact start value loss','Rollout steps','Capped rollout return'],rows,[1.55,1.7,1.65,2.1])
p('Entries are mean ± sample SD over ten training seeds. PI and VI each had 1,000/1,000 successes, mean steps 46.808, and sampled capped return 26.332. Every learned policy also had 1,000/1,000 successes in its evaluation sample. That observation does not imply success probability exactly one, nor does it prove optimality: slower successful episodes lose discounted reward.')
p('Exact 500-step success probabilities computed from the model were at least 0.99988 for every saved greedy learned policy. The planning probability rounds to 1 at floating-point precision. Exact infinite-MDP and capped values agree closely for these final policies; the distinction remains essential for other policies and for learning targets.')
h('Uncertainty and coverage')
for mode in ['fixed','decay']:
 z=diag['comparisons'][mode];lo,hi=z['paired_t_95_interval']
 p(f'{mode.title()} exploration: the paired mean difference in exact start value, SARSA minus Q-learning, is {f(z["sarsa_minus_qlearning_mean"])}; the approximate 95% t interval over ten matched seeds is [{f(lo)}, {f(hi)}]. Both intervals include zero. These small-sample intervals are descriptive and sensitive to unusual seeds.')
p('All 1,404 nonterminal state–action pairs were visited in every run, but visitation was uneven. Some fixed-exploration runs visited their least-sampled pair only once. Decaying exploration improved minimum coverage, but visiting a pair once is far short of accurately estimating its value. Constant α also leaves noise in final estimates.')
p('Exact policy values solve the linear system Vπ=Rπ+γPπVπ with Vπ(g)=0. Local action gaps compare selected actions with Q*. Raw success rates, arrow agreement, and model values answer different questions and are reported separately.')

page('5 Policies produced by the four algorithms')
p(f'The planning maps are identical. The learned maps below use the upper middle run by start-value loss among the ten fixed-exploration seeds: SARSA seed {seeds["sarsa"]} and Q-learning seed {seeds["q_learning"]}. They are representative saved runs, not averages or best-case selections. Rows run from 0 at the top to 19 at the bottom; columns run from 0 on the left to 19 on the right.')
D.add_paragraph().add_run().add_picture(str(ROOT/'figures/policy_comparison.png'),width=Inches(6.7))
p('Orange cells mark an action-value loss above 10⁻⁶ relative to the exact optimal Q values. The shading does not measure how often a state is visited or how large the loss is. A rare-state disagreement can have little start-state effect; a disagreement near a frequently used opening can matter much more.')

page('6 Explaining the observed action differences')
h('Exact exploration reference for this map')
p('SARSA values include the consequences of future exploratory actions. To isolate that effect from finite learning noise, an additional exact model calculation optimizes ε-greedy behavior:')
eq(sub('T','ε'),'V(s) = (1−ε) maxₐ Q(s,a) + (ε/4) ',sub('∑','a∈A'),'Q(s,a)')
p('At ε=0.10, this exact constrained optimum has behavior start value 19.948314550. Removing exploration from its extracted greedy action map gives start value 26.282822039, effectively the unconstrained optimum at the start. Its action map differs from the base optimum only at (14, 0). At ε=0.02, the extracted action map matches the base optimum everywhere.')
table(['Value model at state 14 0','UP action value','RIGHT action value','Preferred'],[('Optimal greedy continuation','31.307577','31.295521','UP'),('Optimal ε=.10 continuation','24.658510','24.672535','RIGHT')],[2.6,1.5,1.5,1.4])
p('This near-tie reversal is a verified exploration effect. It does not support a broad “SARSA takes a safer route” explanation: this environment has no hazard states, and the exact exploratory optimum changes almost none of the greedy decisions.')
h('Concrete disagreements in the representative learned maps')
rows=[]
for st in diag['states']:
 action=['UP','RIGHT','DOWN','LEFT'];rows.append((str(tuple(st['state'])),action[st['reference_action']],action[st['sarsa_action']]+' / '+f(st['sarsa_action_gap']),action[st['q_learning_action']]+' / '+f(st['q_learning_action_gap'])))
table(['State','Optimal action','SARSA action and gap','Q learning action and gap'],rows,[1.0,1.3,2.35,2.35])
p('Gaps use the same reference Q*, not the learners’ own estimates. States were selected from disagreements using optimal-policy discounted occupancy multiplied by the larger action gap, to emphasize relevant locations rather than visually striking rare cells.')
p('At (16, 10), the optimal RIGHT action crosses the opening in the column-10 wall. Its 0.80 intended move reaches (16, 11); both perpendicular slips hit rocks and stay in place, merging into a 0.20 self-loop. The representative Q-learning policy requests LEFT and loses 3.3313 in reference action value, whereas SARSA selects RIGHT. At (18, 8), Q-learning selects the optimal UP action and SARSA selects LEFT, losing 2.7685. Neither method is uniformly better state by state.')
p('The exact exploration reference, the positive start-value losses, uneven counts, and constant step size make finite noisy estimation a stronger explanation for the broad learned-map differences than distinct optimal routes. Maximization bias can affect Q-learning, but this experiment does not isolate it from other learning effects.')

page('7 Learning progress and conclusions')
D.add_paragraph().add_run().add_picture(str(ROOT/'figures/learning_curves.png'),width=Inches(6.9))
p('The curves evaluate the frozen greedy policy exactly at 10,000, 50,000, and 1,000,000 transitions. Error bars show one sample SD across seeds; connecting lines interpolate only those three checkpoints. Early greedy policies can get trapped in recurrent motion, giving value near −100 even while exploratory training continues. The figure does not identify the precise transition at which a useful route was learned.')
h('Why the optimal policy accounts for slips')
p('The best requested action weighs all successor states, including self-loops and detours. It need not point directly toward the goal. Rocks cause delay rather than terminal failure; walls can also suppress an otherwise harmful slip. Bottlenecks, boundaries, and discounting determine whether a requested direction improves expected return. There is no separate collision penalty or hazard penalty.')
h('Supported conclusions and limits')
p('Policy iteration and value iteration agree in both values and actions for this implementation, with negligible numerical error. Learned methods achieve excellent sampled goal rates but do not recover the same optimal values at the tested budget. SARSA with decaying exploration has the highest mean learned start value, while the paired uncertainty intervals do not establish a decisive SARSA versus Q-learning ranking.')
p('The exact ε-greedy reference shows that persistent exploration can change an action ranking, but the effect is localized in this map. Extensive learned-policy disagreements mainly reflect finite estimation and the tested settings. Different seeds, step-size schedules, initialization, and longer budgets could alter the measured comparison.')
p('Ten seeds, one constant learning rate, two exploration schedules, and a single transition budget do not establish asymptotic convergence or a general performance ordering. Further investigation should use denser checkpoints, diminishing per-pair step sizes, additional seeds, and a Double Q-learning ablation if testing maximization bias. Runtime observations are local measurements and should not be treated as universal algorithm speed claims.')

page('8 Reproducibility and sources')
h('Project contents and verification')
p('The RLCompetition2-3-main project includes the unchanged environment and original tests; implementations of all four control methods; saved Q tables, values, policies, and visitation counts; exact and rollout metrics; summary and diagnostic data; policy maps; this editable report; and an AI work log. SOURCE_MANIFEST.json records hashes of the original environment files.')
p('Verification passed 49 tests: 44 original environment tests plus five control tests. The added tests check PI/VI agreement and policy evaluation, merged collisions and goal rewards, fixed-seed training reproducibility, terminal versus nonterminal bootstrap targets, and the planning reference under capped evaluation. The measured runtime used Python '+r['config']['python']+', NumPy '+r['config']['numpy']+', and Gymnasium '+r['config']['gymnasium']+'.')
h('Reproducing the saved experiment')
for text in ['python -m pip install -r requirements.txt','python -m unittest discover -s tests -v','python competition.py --budget 1000000 --seeds 10 --eval-episodes 1000 --output results','python analyze_results.py']:
 pp=p(text)
 for run in pp.runs:run.font.name='Courier New';run.font.size=Pt(9)
p('Run these commands from the project folder. Figure generation uses the optional Pillow package; the saved figures require no additional installation to view. Full installation, smoke-run, and reproduction details are in README.md. Frozen numeric results are in results/results.json and results/summary.csv; diagnostic comparisons are in results/diagnostics.json.')
h('Sources')
refs=[
'[1] HW1 (1).pdf. CIS 5930 Competition Topic 1 Icy Grid World. User-supplied handout, pp. 1–6. The authoritative MDP and Gymnasium contract.',
'[2] RLCompetition1-main original source and Red Team Solution archive. Environment files are byte-identical. Original README and AI log are retained as TEAM_README.md and TEAM_AI_LOG.md.',
'[3] UC Berkeley CS 188. Policy Iteration. https://inst.eecs.berkeley.edu/~cs188/textbook/mdp/policies-iteration.html',
'[4] UC Berkeley CS 188. Value Iteration. https://inst.eecs.berkeley.edu/~cs188/textbook/mdp/value-iteration.html',
'[5] Eric P. Xing. Reinforcement Learning Part 2. Carnegie Mellon University lecture 28, 2008. https://www.cs.cmu.edu/~epxing/Class/10701-08s/Lecture/lecture28-RL2.pdf',
'[6] Farama Foundation. Handling Time Limits. https://gymnasium.farama.org/tutorials/gymnasium_basics/handling_time_limits/',
'[7] Singh, S., Jaakkola, T., Littman, M. L., and Szepesvári, C. Convergence Results for Single-Step On-Policy Reinforcement-Learning Algorithms. Machine Learning, 2000. https://ics.uci.edu/~dechter/courses/ics-295/fall-2019/papers/2000-singh-littmansingh98convergence.pdf']
for text in refs:
 pp=p(text);pp.paragraph_format.space_after=Pt(5)
 for run in pp.runs:run.font.size=Pt(9)
for st in D.styles:
 for el in list(st.element.xpath('.//w:pBdr')):el.getparent().remove(el)
D.save(ROOT/'RLCompetition2-3-main.docx')
print('Report written')
